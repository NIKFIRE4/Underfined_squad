"""Регрессия импорта CSV и выгрузок: синтетические данные, без БД и модели."""
import csv
import io
import json
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from openpyxl import load_workbook

from exports import export_lot, xlsx_bytes
from pipeline import read_csv


@pytest.fixture
def job_folder():
    # Каждый тест получает собственные файлы; удаление выполняется и при падении.
    with TemporaryDirectory(prefix="procurement-tests-") as directory:
        yield Path(directory)


@pytest.fixture
def csv_file(job_folder):
    def write(rows, encoding="utf-8-sig", delimiter=";"):
        path = job_folder / "input.csv"
        with path.open("w", encoding=encoding, newline="") as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            writer.writerows(rows)
        return path
    return write


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1251"])
@pytest.mark.parametrize("delimiter", [";", ","])
def test_csv_formats_preserve_quoted_subject_and_identifiers(csv_file, encoding, delimiter):
    """IMP-01: поддерживаемые CSV сохраняют кириллицу, кавычки и ведущие нули."""
    subject = 'Бумага; картон, "А4"\nВторая строка'
    path = csv_file([["lot_id", "subject"], ["000042", subject]], encoding, delimiter)
    records = list(read_csv(path, "notices"))
    assert len(records) == 1
    assert records[0][0] == {"lot_id": "000042", "subject": subject}
    assert records[0][3] == delimiter


@pytest.mark.parametrize("headers", [
    ["lot_id", "subject", "SUBJECT"],
    ["lot_id", "subject", " subject "],
    ["lot_id", "", "subject"],
], ids=["case-duplicate", "whitespace-duplicate", "empty-header"])
def test_csv_rejects_ambiguous_headers(csv_file, headers):
    """IMP-02: неоднозначные заголовки не должны молча терять данные."""
    path = csv_file([headers, ["000042", "Бумага", "Картон"]])
    with pytest.raises(ValueError, match="пустые или повторяются"):
        list(read_csv(path, "notices"))


@pytest.mark.parametrize("bad_row, message", [
    ([" ", "Бумага"], "пустой lot_id"),
    (["000042", "  "], "пустой предмет"),
    (["000042"], "число значений"),
    (["000042", "Бумага", "лишнее"], "число значений"),
])
def test_csv_rejects_incomplete_or_extra_data(csv_file, bad_row, message):
    """IMP-03: некорректная запись отклоняется с понятной причиной."""
    with pytest.raises(ValueError, match=message):
        list(read_csv(csv_file([["lot_id", "subject"], bad_row]), "notices"))


def test_utf8_character_crossing_encoding_sample_boundary(job_folder):
    """IMP-04: срез 64 KiB внутри символа UTF-8 не меняет кодировку файла."""
    header = b"lot_id;subject\n000042;"
    subject = "A" * (65535 - len(header)) + "Я" + "Бумага" * 100
    path = job_folder / "large.csv"
    path.write_bytes(header + subject.encode("utf-8") + b"\n")
    records = list(read_csv(path, "notices"))
    assert records[0][0]["subject"] == subject
    assert records[0][2] == "utf-8-sig"


@pytest.fixture
def saved_result(job_folder):
    fields = ["lot_id", "supplier_inn", "supplier_kpp", "supplier_name", "is_new"]
    selected = [dict(zip(fields, [
        "000042", f"{index:010d}", "001234567",
        f'Компания {index}; "Север"\nФилиал', str(bool(index % 2)),
    ])) for index in range(151)]
    other = dict(zip(fields, ["000043", "9999999999", "", "Другой лот", "False"]))
    with (job_folder / "suppliers.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fields, delimiter=";")
        writer.writeheader()
        # Чужой лот между нужными строками ловит выбор непрерывного среза.
        writer.writerows([*selected[:70], other, *selected[70:]])
    (job_folder / "lots.jsonl").write_text(
        json.dumps({"lot_id": "empty"}) + "\n", encoding="utf-8"
    )
    return job_folder, fields, selected


@contextmanager
def workbook_from(payload):
    with io.BytesIO(payload) as stream:
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            yield book
        finally:
            book.close()


def exported_rows(payload, file_format):
    if file_format == "csv":
        return list(csv.reader(io.StringIO(payload.decode("utf-8-sig")), delimiter=";"))
    with workbook_from(payload) as book:
        return [["" if cell is None else cell for cell in row]
                for row in book.active.iter_rows(values_only=True)]


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_lot_export_includes_more_than_preview_limit(saved_result, file_format):
    """EXP-01: экспорт содержит все 151 строку лота, а не только preview."""
    folder, fields, selected = saved_result
    payload, mime = export_lot(folder, "000042", file_format)
    assert mime == {"csv": "text/csv; charset=utf-8",
                    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}[file_format]
    assert exported_rows(payload, file_format) == [fields] + [[r[key] for key in fields] for r in selected]


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_empty_known_lot_exports_headers(saved_result, file_format):
    """EXP-02: известный лот без поставщиков даёт файл с заголовками."""
    folder, fields, _ = saved_result
    payload, _ = export_lot(folder, "empty", file_format)
    assert exported_rows(payload, file_format) == [fields]


@pytest.mark.parametrize("length", [32766, 32767, 32768])
def test_excel_cell_length_boundary(length):
    """EXP-03: предел Excel 32767 символов; превышение не обрезается молча."""
    value = "Я" * length
    if length > 32767:
        with pytest.raises(ValueError, match="лимит ячейки"):
            xlsx_bytes(["name"], [{"name": value}])
    else:
        payload = xlsx_bytes(["name"], [{"name": value}])
        assert exported_rows(payload, "xlsx") == [["name"], [value]]


@pytest.mark.parametrize("value", ["=1+1", "+SUM(1,2)", "@SUM(1)", "-1+2"])
def test_xlsx_values_are_text_not_executable_formulas(value):
    """EXP-04: похожее на формулу имя остаётся строкой, а не формулой Excel."""
    payload = xlsx_bytes(["name"], [{"name": value}])
    with workbook_from(payload) as book:
        cell = book.active["A2"]
        assert cell.value == value
        assert cell.data_type == "s"
